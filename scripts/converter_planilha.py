"""Converte a planilha antiga "CHECK LIST RASTREADOR RASTER" em eventos do painel.

Uso:
  python3 scripts/converter_planilha.py PLANILHA.xlsx --revisar
      lista as correções propostas (placas e datas), sem gerar nada
  python3 scripts/converter_planilha.py PLANILHA.xlsx --gerar SAIDA_DIR [--decisoes decisoes.json]
      grava SAIDA_DIR/eventos/*.json (um veiculo_cadastrado por placa)

Datas: a planilha mistura texto "dd/mm/aaaa" (confiável) com datas que o Excel leu com dia e
mês trocados. Para cada data com dia <= 12, as duas leituras são testadas contra a janela
plausível da situação; quando só uma cabe, ela é usada; quando as duas ou nenhuma cabem, a
linha vai para decisão manual (arquivo --decisoes, {"PLACA": "dd/mm/aaaa"}).
"""
import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

HOJE = dt.date(2026, 10, 15)
VALIDADE_MAX = 31  # dias: um checklist aprovado nunca vence mais de ~30 dias depois do teste

SITUACOES = {
    'VALIDADO E LANÇADO': 'validado',
    'MANUTENÇÃO': 'manutencao',
    'CHECK LIST CANCELADO': 'cancelado',
    'CHECK LIST FEITO CONFIRMAR VALIDAÇÃO': 'confirmar',
    'NECESSÁRIO REALIZAÇÃO CHECK LIST': 'necessario',
}
# erros de digitação conhecidos (confirmados pelo PDF/portal: a placa real é UVR...)
PLACAS_CORRIGIDAS = {'URV1I58': 'UVR1I58', 'URV1I53': 'UVR1I53', 'URV1I57': 'UVR1I57', 'URV1I63': 'UVR1I63'}
PLACA_OK = re.compile(r'^[A-Z]{3}\d[A-Z0-9]\d{2}$')


def normalizar_placa(p):
    return re.sub(r'[^A-Z0-9]', '', str(p or '').upper())


def janela(situacao):
    """Intervalo em que o vencimento é plausível para a situação."""
    if situacao in ('validado', 'confirmar'):
        # checklist em dia na planilha: vence no máximo ~30 dias à frente, e no máximo
        # ~30 dias atrás se a planilha não foi atualizada a tempo
        return HOJE - dt.timedelta(days=VALIDADE_MAX), HOJE + dt.timedelta(days=VALIDADE_MAX)
    if situacao == 'necessario':
        return HOJE - dt.timedelta(days=120), HOJE + dt.timedelta(days=VALIDADE_MAX)
    return None  # manutenção/cancelado: pode estar vencido há muito tempo, sem janela segura


def ler_data(valor, situacao):
    """Devolve (data, origem, alternativa) ou (None, motivo, candidatos)."""
    if isinstance(valor, str):
        m = re.fullmatch(r'\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*', valor)
        if not m:
            return None, f'texto não reconhecido: {valor!r}', []
        d, mo, a = map(int, m.groups())
        return dt.date(a, mo, d), 'texto', None
    if isinstance(valor, dt.datetime):
        lida = valor.date()
        if lida.day > 12:
            return lida, 'data', None  # sem ambiguidade
        trocada = dt.date(lida.year, lida.day, lida.month)
        if trocada == lida:
            return lida, 'data', None
        jan = janela(situacao)
        if jan:
            cabe = [c for c in (lida, trocada) if jan[0] <= c <= jan[1]]
            if len(cabe) == 1:
                return cabe[0], ('data trocada' if cabe[0] == trocada else 'data'), (trocada if cabe[0] == lida else lida)
        return None, 'ambígua', [lida, trocada]
    return None, 'sem data', []


def ler_planilha(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True).worksheets[0]
    linhas = []
    for r in range(10, ws.max_row + 1):
        placa_bruta = ws.cell(r, 3).value
        if not placa_bruta:
            continue
        status = str(ws.cell(r, 5).value or '').strip()
        obs = str(ws.cell(r, 7).value or '').strip()
        mot = str(ws.cell(r, 6).value or '').strip()
        linhas.append({'linha': r, 'placa_bruta': str(placa_bruta).strip(), 'vencimento_bruto': ws.cell(r, 4).value,
                       'status': status, 'observacao': ' · '.join(x for x in (obs, f'Motorista: {mot}' if mot else '') if x)})
    return linhas


def analisar(caminho, decisoes=None):
    decisoes = decisoes or {}
    itens, correcoes, pendentes = [], [], []
    vistas = {}
    for l in ler_planilha(caminho):
        placa = normalizar_placa(l['placa_bruta'])
        if placa in PLACAS_CORRIGIDAS:
            correcoes.append({'linha': l['linha'], 'tipo': 'placa', 'de': placa, 'para': PLACAS_CORRIGIDAS[placa]})
            placa = PLACAS_CORRIGIDAS[placa]
        if not PLACA_OK.match(placa):
            pendentes.append({'linha': l['linha'], 'placa': placa, 'motivo': 'placa fora do padrão'})
        sit = SITUACOES.get(l['status'])
        if not sit:
            pendentes.append({'linha': l['linha'], 'placa': placa, 'motivo': f'status desconhecido: {l["status"]!r}'})
            continue
        if placa in decisoes:
            venc = dt.datetime.strptime(decisoes[placa], '%d/%m/%Y').date()
            origem = 'decisão'
        else:
            venc, origem, extra = ler_data(l['vencimento_bruto'], sit)
            if venc is None:
                pendentes.append({'linha': l['linha'], 'placa': placa, 'situacao': sit, 'motivo': f'vencimento {origem}',
                                  'opcoes': [c.strftime('%d/%m/%Y') for c in extra]})
                continue
            if origem == 'data trocada':
                correcoes.append({'linha': l['linha'], 'tipo': 'data', 'placa': placa,
                                  'de': extra.strftime('%d/%m/%Y'), 'para': venc.strftime('%d/%m/%Y')})
        if placa in vistas:
            pendentes.append({'linha': l['linha'], 'placa': placa, 'motivo': f'placa repetida (também na linha {vistas[placa]})'})
            continue
        vistas[placa] = l['linha']
        itens.append({'linha': l['linha'], 'placa': placa, 'situacao': sit, 'vencimento': venc, 'observacao': l['observacao'], 'origem': origem})
    return itens, correcoes, pendentes


def evento(item, ts):
    id_ = 'planilha-' + hashlib.sha1(item['placa'].encode()).hexdigest()[:12]  # fixo: extrair de novo não duplica
    dados = {'vencimento': item['vencimento'].strftime('%Y-%m-%dT00:00:00'), 'situacao': item['situacao'], 'linhaPlanilha': item['linha']}
    if item['observacao']:
        dados['observacao'] = item['observacao']
    return id_, {'id': id_, 'ts': ts, 'pessoa': 'Planilha antiga', 'pc': 'conversao', 'tipo': 'veiculo_cadastrado', 'placa': item['placa'], 'dados': dados}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('planilha')
    ap.add_argument('--revisar', action='store_true')
    ap.add_argument('--gerar')
    ap.add_argument('--decisoes')
    a = ap.parse_args()
    dec = json.loads(Path(a.decisoes).read_text()) if a.decisoes else {}
    itens, correcoes, pendentes = analisar(a.planilha, dec)
    if a.revisar or not a.gerar:
        print(json.dumps({'itens': len(itens), 'correcoes': correcoes, 'pendentes': pendentes,
                          'por_situacao': {s: sum(1 for i in itens if i['situacao'] == s) for s in SITUACOES.values()},
                          'linhas': [{**i, 'vencimento': i['vencimento'].strftime('%d/%m/%Y')} for i in itens]},
                         ensure_ascii=False, indent=1))
        return
    if pendentes:
        sys.exit(f'{len(pendentes)} linha(s) pendente(s): rode --revisar e informe --decisoes.')
    saida = Path(a.gerar) / 'eventos'
    saida.mkdir(parents=True, exist_ok=True)
    ts = '2026-10-15T00:00:00.000Z'  # antes de qualquer alteração feita no painel
    for item in itens:
        id_, ev = evento(item, ts)
        (saida / f'2026-10-15T00-00-00-000Z_planilha-antiga_{id_}.json').write_text(json.dumps(ev, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'{len(itens)} eventos gravados em {saida}')


if __name__ == '__main__':
    main()
